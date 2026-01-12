#!/usr/bin/env python
"""Convert an OCR'd PDF into an editable PPTX with Traditional Chinese OCR for images."""

import io
import sys
from pathlib import Path

import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt


def extract_page_text(pdf_path: Path) -> list[str]:
    texts: list[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            texts.append(text.strip())
    return texts


def ocr_image(image_bytes: bytes) -> str:
    image = Image.open(io.BytesIO(image_bytes))
    return pytesseract.image_to_string(image, lang="chi_tra").strip()


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float) -> None:
    textbox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    text_frame = textbox.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(16)


def build_pptx(pdf_path: Path, output_path: Path) -> None:
    presentation = Presentation()
    page_texts = extract_page_text(pdf_path)

    doc = fitz.open(str(pdf_path))

    for page_index, page in enumerate(doc):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        page_text = page_texts[page_index] if page_index < len(page_texts) else ""
        if page_text:
            add_textbox(slide, page_text, left=0.5, top=0.3, width=12.3, height=3.5)

        images = page.get_images(full=True)
        if images:
            current_top = 4.0
            for image_index, image_info in enumerate(images):
                xref = image_info[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]

                image_stream = io.BytesIO(image_bytes)
                slide.shapes.add_picture(image_stream, Inches(0.5), Inches(current_top), width=Inches(5.5))

                ocr_text = ocr_image(image_bytes)
                if ocr_text:
                    add_textbox(
                        slide,
                        f"圖片 {image_index + 1} OCR:\n{ocr_text}",
                        left=6.2,
                        top=current_top,
                        width=6.0,
                        height=2.0,
                    )

                current_top += 2.5

    presentation.save(str(output_path))


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: python pdf_to_pptx.py <input_ocr.pdf> <output.pptx>")
        return 1

    input_pdf = Path(sys.argv[1]).resolve()
    output_pptx = Path(sys.argv[2]).resolve()

    if not input_pdf.exists():
        print(f"Input PDF not found: {input_pdf}")
        return 1

    build_pptx(input_pdf, output_pptx)
    print(f"Saved PPTX: {output_pptx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
